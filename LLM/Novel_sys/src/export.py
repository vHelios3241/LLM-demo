# uv add python-docx

from __future__ import annotations

import re
from pathlib import Path
from typing import Union

from LLM.src.schema import Novel


def _sanitize_filename(name: str) -> str:
    """将字符串中的非法文件名字符替换为下划线。"""
    return re.sub(r'[\\/:*?"<>|]', '_', name)


def export_markdown(novel: Novel, output_dir: Union[str, Path]) -> Path:
    """将小说导出为 Markdown 文件。

    Args:
        novel: 已填充各章正文的小说对象。
        output_dir: 输出目录路径，不存在时会自动创建。

    Returns:
        生成的 .md 文件路径。
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    lines.append(f"# 《{novel.title}》")
    lines.append("")
    lines.append(novel.full_summary)
    lines.append("")

    for chapter in novel.chapters:
        lines.append(f"## 第{chapter.order}章 {chapter.title}")
        lines.append("")
        if chapter.body:
            lines.append(chapter.body)
        else:
            lines.append("")
        lines.append("")

    safe_name = _sanitize_filename(novel.title)
    md_path = output_dir / f"{safe_name}.md"
    md_path.write_text("\n".join(lines), encoding="utf-8")

    return md_path


def export_word(novel: Novel, output_dir: Union[str, Path]) -> Path:
    """将小说导出为 Word (.docx) 文件，每章分页。

    注意：依赖 python-docx 库，如未安装请执行 ``uv add python-docx``。

    Args:
        novel: 已填充各章正文的小说对象。
        output_dir: 输出目录路径，不存在时会自动创建。

    Returns:
        生成的 .docx 文件路径。
    """
    # 延迟导入，确保没装 python-docx 时不影响 Markdown 导出
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    document = Document()

    # --- 书名 ---
    heading = document.add_heading(f"《{novel.title}》", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- 梗概 ---
    document.add_paragraph(novel.full_summary)

    # --- 各章 ---
    for i, chapter in enumerate(novel.chapters):
        document.add_heading(f"第{chapter.order}章 {chapter.title}", level=1)

        if chapter.body:
            document.add_paragraph(chapter.body)
        else:
            document.add_paragraph("")

        # 最后一章之后不加分页
        if i < len(novel.chapters) - 1:
            document.add_page_break()

    safe_name = _sanitize_filename(novel.title)
    docx_path = output_dir / f"{safe_name}.docx"
    document.save(str(docx_path))

    return docx_path